"""Gera o roteiro de fala (partes/01_roteiro.md) e o cronograma de 00_intro.md a partir das notas do PPTX.

Uso: python3 gerar_roteiro.py   (chamado também por montar_guia.py)
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
PPTX = HERE.parent / "apresentacao" / "SEMICA_Fumaca_Cirurgica_Baccaro.pptx"
PARTES = HERE / "partes"
EMU_IN = 914400


def titulo(slide, i, n):
    if i == 1:
        return "Capa"
    if i == n:
        return "Encerramento"
    tops = sorted((sh.top, sh.text_frame.text.strip()) for sh in slide.shapes
                  if sh.has_text_frame and sh.text_frame.text.strip() and sh.top < 1.9 * EMU_IN)
    partes = [t for _, t in tops][:2]
    return ": ".join(partes) if partes else f"Slide {i}"


def mmss(s):
    return f"{s // 60}:{s % 60:02d}"


def main():
    prs = Presentation(str(PPTX))
    n = len(prs.slides)
    md = ["# Roteiro de fala, slide a slide (notas do orador)\n"]
    linhas_crono = []
    fim = 0
    blocos = []
    for i, s in enumerate(prs.slides, start=1):
        notas = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        blocos.append(f"## Slide {i} — {titulo(s, i, n)}\n\n{notas}\n")
        m = re.match(r"\[(\d+):(\d+)–(\d+):(\d+)\]\s*(.+)", notas)
        if m:
            a = int(m.group(1)) * 60 + int(m.group(2))
            b = int(m.group(3)) * 60 + int(m.group(4))
            secao = re.sub(r"\s*\(~\d+ s\).*$", "", m.group(5)).split("←")[0].strip()
            linhas_crono.append(f"| {i} | {secao.capitalize().replace("grade", "GRADE") if secao.isupper() else secao} | {mmss(b - a)} | {mmss(b)} |")
            fim = max(fim, b)
        elif i == n:
            linhas_crono.append(f"| {i} | Encerramento (fica na tela, com o QR code, durante as perguntas) | — | — |")
    md.append(f"Tempo-alvo total: ~{fim // 60} min {fim % 60:02d} s, deixando folga para os 10 minutos. "
              "Os colchetes indicam o relógio acumulado.\n")
    md.extend(blocos)
    (PARTES / "01_roteiro.md").write_text("\n".join(md), encoding="utf-8")

    intro = (PARTES / "00_intro.md").read_text(encoding="utf-8")
    tabela = "\n".join(["| Slide | Seção | Tempo | Relógio acumulado |", "|---|---|---|---|", *linhas_crono])
    intro = re.sub(r"(## Cronograma de fala \(tempo-alvo por slide\)\n\n)\| Slide.*?\n\n", lambda m: m.group(1) + tabela + "\n\n",
                   intro, count=1, flags=re.S)
    (PARTES / "00_intro.md").write_text(intro, encoding="utf-8")
    print(f"OK: roteiro de {n} slides, {mmss(fim)}")


if __name__ == "__main__":
    main()
