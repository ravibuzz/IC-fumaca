"""Gera o CADERNO DO ORADOR da versão enxuta em PDF: cada slide com a miniatura, o tempo, a fala e as notas de apoio.

Uso:  python3 gerar_caderno.py   (depois de build_deck_enxuto.py e da conversão do PPTX em PDF)

Lê as notas do PPTX e as imagens do PDF da apresentação; monta um HTML e imprime em PDF com o Chromium (headless).
"""
from __future__ import annotations

import base64
import html
import re
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
PPTX = HERE / "SEMICA_Fumaca_Cirurgica_Baccaro_ENXUTA.pptx"
PDF_SLIDES = HERE / "SEMICA_Fumaca_Cirurgica_Baccaro_ENXUTA.pdf"
ROTEIRO_MD = HERE / "ROTEIRO_VERSAO_ENXUTA.md"
OUT = HERE / "CADERNO_DO_ORADOR_VERSAO_ENXUTA.pdf"
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

CSS = """
@page { size: A4; margin: 14mm 14mm 16mm 14mm; }
body { font-family: "Liberation Sans", Arial, sans-serif; color: #1f2430; font-size: 10.5pt; line-height: 1.38; }
h1 { color: #002060; font-size: 20pt; margin: 0 0 4px; }
h2 { color: #002060; font-size: 13.5pt; margin: 0 0 6px; border-bottom: 2px solid #C9A227; padding-bottom: 3px; }
.sub { color: #555; margin: 0 0 14px; }
table { border-collapse: collapse; width: 100%; margin: 8px 0 14px; font-size: 10pt; }
th { background: #002060; color: #fff; text-align: left; padding: 5px 7px; }
td { border-bottom: 1px solid #d8dce5; padding: 4px 7px; }
.box { background: #f3f5f9; border-left: 4px solid #002060; padding: 8px 12px; margin: 8px 0 12px; }
.slide { page-break-inside: avoid; break-inside: avoid; margin-bottom: 16px; }
.slide.novo { page-break-before: always; break-before: page; }
.slide img { width: 100%; border: 1px solid #c9cfda; border-radius: 4px; margin: 2px 0 8px; }
.fala { background: #fff8e1; border-left: 4px solid #C9A227; padding: 8px 12px; margin: 6px 0 8px; font-size: 11.5pt; }
.fala b { color: #8A6D1F; }
.rot { color: #002060; font-weight: bold; margin: 8px 0 2px; }
ul { margin: 2px 0 6px 18px; padding: 0; }
li { margin: 1px 0; }
p { margin: 3px 0; }
.tag { display: inline-block; background: #002060; color: #fff; border-radius: 3px; padding: 1px 7px; font-size: 9.5pt;
       margin-right: 6px; }
.apoio .tag { background: #8A6D1F; }
"""


def slide_images(tmp):
    subprocess.run(["pdftoppm", "-r", "60", "-png", str(PDF_SLIDES), str(Path(tmp) / "s")], check=True)
    return sorted(Path(tmp).glob("s-*.png"))


def nota_html(text):
    """Primeira linha = título; 'FALA:' vira a caixa da fala; '- ' vira lista; linhas terminadas em ':' viram rótulos."""
    lines = text.split("\n")
    titulo, corpo = lines[0], lines[1:]
    out, lista, modo_fala, fala = [], [], False, []

    def fecha_lista():
        if lista:
            out.append("<ul>" + "".join(f"<li>{x}</li>" for x in lista) + "</ul>")
            lista.clear()

    def fecha_fala():
        nonlocal modo_fala
        if modo_fala:
            out.append('<div class="fala"><b>FALA</b><br>' + "<br>".join(fala) + "</div>")
            fala.clear()
            modo_fala = False

    for ln in corpo:
        s = ln.strip()
        if s == "FALA:":
            fecha_lista()
            modo_fala = True
            continue
        if modo_fala:
            if s:
                fala.append(html.escape(s))
                continue
            fecha_fala()
            continue
        if not s:
            fecha_lista()
            continue
        if s.startswith("- "):
            lista.append(html.escape(s[2:]))
        elif s.endswith(":") and len(s) < 90:
            fecha_lista()
            out.append(f'<div class="rot">{html.escape(s[:-1])}</div>')
        else:
            fecha_lista()
            out.append(f"<p>{html.escape(s)}</p>")
    fecha_fala()
    fecha_lista()
    return titulo, "".join(out)


def tabela_tempos():
    rows = [ln for ln in ROTEIRO_MD.read_text(encoding="utf-8").split("\n") if re.match(r"\| \d+ \|", ln)]
    corpo = "".join("<tr>" + "".join(f"<td>{html.escape(c.strip())}</td>" for c in r.strip("|").split("|")) + "</tr>"
                    for r in rows)
    return ("<table><tr><th>Slide</th><th>Seção</th><th>Tempo</th><th>Palavras</th><th>Relógio</th></tr>"
            + corpo + "</table>")


def main():
    prs = Presentation(str(PPTX))
    with tempfile.TemporaryDirectory() as tmp:
        imgs = slide_images(tmp)
        assert len(imgs) == len(prs.slides), (len(imgs), len(prs.slides))
        partes = [
            "<h1>Caderno do orador — versão enxuta</h1>",
            '<p class="sub">VIII SEMICA · 21/10/2026 · Theo Báccaro · Revisão sistemática e metanálise de prevalência: '
            "fumaça cirúrgica e saúde da equipe de centro cirúrgico</p>",
            '<div class="box"><b>Como usar.</b> Cada slide traz a <b>fala</b> (o que dizer em voz alta, em amarelo) e o '
            "<b>apoio</b> (números e respostas prontas, só para consultar se perguntarem). Tempo-alvo: <b>7:40</b> "
            "(limite de 10 min). Se estiver mais de 15 s atrasado num marco do relógio, corte o detalhe do slide "
            "seguinte e diga só a frase-chave (o subtítulo). Os slides 16 a 24 são de apoio: no modo de apresentação, "
            "digite o número e Enter para ir direto; 15 e Enter para voltar ao encerramento.</div>",
            "<h2>Tempo por slide</h2>",
            tabela_tempos(),
        ]
        for i, (sl, img) in enumerate(zip(prs.slides, imgs), start=1):
            titulo, corpo = nota_html(sl.notes_slide.notes_text_frame.text)
            apoio = i > 15
            b64 = base64.b64encode(img.read_bytes()).decode()
            cls = "slide novo" + (" apoio" if apoio else "")
            tag = f"Apoio {i}" if apoio else f"Slide {i}"
            partes.append(f'<div class="{cls}"><h2><span class="tag">{tag}</span>{html.escape(titulo)}</h2>'
                          f'<img src="data:image/png;base64,{b64}">{corpo}</div>')
        page = Path(tmp) / "caderno.html"
        page.write_text(f"<!doctype html><html lang='pt-BR'><head><meta charset='utf-8'><style>{CSS}</style></head>"
                        f"<body>{''.join(partes)}</body></html>", encoding="utf-8")
        subprocess.run([CHROME, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={OUT}", page.as_uri()], check=True, capture_output=True)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
