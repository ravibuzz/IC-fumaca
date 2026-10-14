"""Helpers para montar os slides sobre o modelo oficial do SEMICA (python-pptx).

Canvas do modelo: 20,0" x 11,25" (18288000 x 10287000 EMU).
"""
from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
ICONS = ASSETS / "icones"

# ---------- paleta (institucional: azul-marinho + dourado do modelo) ----------
NAVY = RGBColor(0x00, 0x20, 0x60)
NAVY_DARK = RGBColor(0x1A, 0x01, 0x5F)
GOLD = RGBColor(0xC8, 0xA9, 0x51)
RED = RGBColor(0x8B, 0x00, 0x00)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xE0, 0x8A, 0x1E)
INK = RGBColor(0x1F, 0x24, 0x33)
GRAY = RGBColor(0x5A, 0x62, 0x70)
LIGHT = RGBColor(0xF2, 0xF4, 0xF8)
NAVY_TINT = RGBColor(0xE6, 0xEB, 0xF5)
GOLD_TINT = RGBColor(0xFB, 0xF5, 0xE3)
RED_TINT = RGBColor(0xF9, 0xE9, 0xE9)
TEAL_TINT = RGBColor(0xE3, 0xF3, 0xF4)
GREEN_TINT = RGBColor(0xE7, 0xF3, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Arial"

# ---------- geometria do modelo (polegadas) ----------
SLIDE_W, SLIDE_H = 20.0, 11.25
L, R = 1.25, 19.45          # margens de conteúdo
TOP = 1.85                   # início do conteúdo
BOT = 10.75                  # fim do conteúdo (sem trilha de passos)
BOT_T = 10.0                 # fim do conteúdo quando há trilha de passos

STEPS = ["Pergunta e protocolo", "Busca", "Seleção", "Extração",
         "Risco de viés", "Síntese", "Certeza (GRADE)"]


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


# ---------- texto ----------
def _apply_run(run, size=None, bold=None, italic=None, color=None, font=FONT):
    f = run.font
    f.name = font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def _set_bullet(p, char="•", indent=0.3, color=None, size_pct=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent))))
    pPr.set("indent", str(-Emu(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buSzPct", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if color is not None:
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(buClr, qn("a:srgbClr"))
        srgb.set("val", str(color))
    if size_pct is not None:
        buSz = etree.SubElement(pPr, qn("a:buSzPct"))
        buSz.set("val", str(int(size_pct * 1000)))
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", char)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def fill_text(tf, paragraphs, size=22, color=INK, bold=False, align=PP_ALIGN.LEFT,
              space_after=6, line_spacing=1.0, bullet=None, bullet_color=None, font=FONT):
    """paragraphs: lista; cada item é str ou dict(runs=[(texto, opts)...], **par_opts)
    ou lista de runs [(texto, opts), ...]. opts: size,bold,italic,color."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    first = True
    for item in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        par_opts = {}
        runs = item
        if isinstance(item, dict):
            par_opts = {k: v for k, v in item.items() if k != "runs"}
            runs = item["runs"]
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = par_opts.get("align", align)
        p.space_after = Pt(par_opts.get("space_after", space_after))
        if par_opts.get("space_before") is not None:
            p.space_before = Pt(par_opts["space_before"])
        p.line_spacing = par_opts.get("line_spacing", line_spacing)
        for text, opts in runs:
            r = p.add_run()
            r.text = text
            _apply_run(r, size=opts.get("size", par_opts.get("size", size)),
                       bold=opts.get("bold", par_opts.get("bold", bold)),
                       italic=opts.get("italic", par_opts.get("italic", False)),
                       color=opts.get("color", par_opts.get("color", color)), font=font)
        b = par_opts.get("bullet", bullet)
        if b:
            _set_bullet(p, char=b if isinstance(b, str) else "•", indent=par_opts.get("indent", 0.3),
                        color=par_opts.get("bullet_color", bullet_color))
        else:
            _no_bullet(p)
    return tf


def tb(slide, x, y, w, h, paragraphs, size=22, color=INK, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, margin=0.0, space_after=6, line_spacing=1.0, bullet=None,
       bullet_color=None, font=FONT, name=None, rotation=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    if rotation is not None:
        box.rotation = rotation
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    fill_text(tf, paragraphs, size=size, color=color, bold=bold, align=align, space_after=space_after,
              line_spacing=line_spacing, bullet=bullet, bullet_color=bullet_color, font=font)
    return box


# ---------- formas ----------
def _strip_style(shp):
    """Remove <p:style> (sombra/estilo herdado do tema) para render idêntico no PowerPoint e no LibreOffice."""
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)


def rect(slide, x, y, w, h, fill=LIGHT, line=None, line_w=1.0, radius=0.06, shadow=False, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(shp)
    if radius:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    if not shadow:
        shp.shadow.inherit = False
    if name:
        shp.name = name
    # sem texto por padrão
    shp.text_frame.text = ""
    return shp


def circle(slide, cx, cy, d, fill=NAVY, line=None, text=None, size=22, color=WHITE, bold=True, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    _strip_style(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    if name:
        shp.name = name
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if text is not None:
        fill_text(tf, [text], size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, space_after=0)
    else:
        tf.text = ""
    return shp


def line(slide, x1, y1, x2, y2, color=NAVY, width=2.0, arrow=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _strip_style(c)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        prstDash = etree.SubElement(ln, qn("a:prstDash"))
        prstDash.set("val", "dash")
    if arrow:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def chevron(slide, x, y, w, h, fill=NAVY, text=None, size=18, color=WHITE, bold=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if text:
        fill_text(tf, [text], size=size, color=color, bold=bold, align=PP_ALIGN.CENTER, space_after=0)
    else:
        tf.text = ""
    return shp


# ---------- imagens ----------
def pic(slide, path, x, y, w=None, h=None, name=None):
    """Insere imagem mantendo proporção dentro de (w, h). Retorna a forma."""
    path = str(path)
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    if w is not None and h is not None:
        if w / h > ar:
            pw, ph = h * ar, h
        else:
            pw, ph = w, w / ar
        px = x + (w - pw) / 2
        py = y + (h - ph) / 2
    elif w is not None:
        pw, ph, px, py = w, w / ar, x, y
    else:
        pw, ph, px, py = h * ar, h, x, y
    shp = slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))
    if name:
        shp.name = name
    return shp


def icon(slide, name, x, y, size=0.5, color="navy"):
    return pic(slide, ICONS / f"{name}_{color}.png", x, y, size, size, name=f"icone {name}")


def icon_badge(slide, name, cx, cy, d=0.9, fill=NAVY, color="white", icon_frac=0.5):
    circle(slide, cx, cy, d, fill=fill)
    s = d * icon_frac
    icon(slide, name, cx - s / 2, cy - s / 2, s, color=color)


# ---------- componentes compostos ----------
def card(slide, x, y, w, h, title=None, body=None, fill=LIGHT, title_color=NAVY, body_size=20,
         title_size=22, icon_name=None, icon_color="white", badge_fill=NAVY, pad=0.3, line_color=None,
         bullet=None, body_color=INK, title_bold=True, radius=0.06, body_space_after=5):
    rect(slide, x, y, w, h, fill=fill, line=line_color, radius=radius)
    ty = y + pad
    tx = x + pad
    if icon_name:
        d = 0.8
        icon_badge(slide, icon_name, x + pad + d / 2, y + pad + d / 2, d=d, fill=badge_fill, color=icon_color)
        tx = x + pad + d + 0.25
    if title:
        tb(slide, tx, ty, w - (tx - x) - pad, 0.8 if icon_name else 0.5, [title], size=title_size,
           color=title_color, bold=title_bold, anchor=MSO_ANCHOR.MIDDLE if icon_name else MSO_ANCHOR.TOP,
           space_after=0)
        ty += (0.95 if icon_name else 0.55)
    if body:
        tb(slide, x + pad, ty, w - 2 * pad, y + h - ty - pad * 0.8, body, size=body_size, color=body_color,
           bullet=bullet, space_after=body_space_after)
    return ty


def stat(slide, x, y, w, number, label, h=1.6, num_size=54, lab_size=18, num_color=NAVY, lab_color=GRAY,
         fill=None, align=PP_ALIGN.LEFT, number_h=None):
    if fill is not None:
        rect(slide, x, y, w, h, fill=fill)
    pad = 0.2 if fill is not None else 0.0
    nh = number_h or (num_size / 72 * 1.25)
    tb(slide, x + pad, y + pad * 0.6, w - 2 * pad, nh, [number], size=num_size, color=num_color, bold=True,
       align=align, space_after=0, line_spacing=0.95)
    tb(slide, x + pad, y + pad * 0.6 + nh + 0.02, w - 2 * pad, h - nh - pad, [label], size=lab_size,
       color=lab_color, align=align, space_after=0, line_spacing=1.05)


def step_track(slide, active, y=10.12):
    """Trilha dos 7 passos no rodapé. active: int ou conjunto de ints (1..7)."""
    if isinstance(active, int):
        active = {active}
    n = len(STEPS)
    x0, x1 = L + 0.1, R - 0.1
    seg = (x1 - x0) / n
    line(slide, x0 + seg / 2, y + 0.22, x1 - seg / 2, y + 0.22, color=rgb("C9CFDA"), width=2.5)
    for i, name in enumerate(STEPS, start=1):
        cx = x0 + seg * (i - 0.5)
        is_active = i in active
        circle(slide, cx, y + 0.22, 0.44 if is_active else 0.36, fill=RED if is_active else rgb("C9CFDA"),
               text=str(i), size=16 if is_active else 13, color=WHITE)
        tb(slide, cx - seg / 2, y + 0.5, seg, 0.4, [name], size=14, color=NAVY if is_active else GRAY,
           bold=is_active, align=PP_ALIGN.CENTER, space_after=0)


# ---------- tabela ----------
def table(slide, x, y, w, rows, col_widths=None, header=True, size=16, row_h=0.42, header_fill=NAVY,
          header_color=WHITE, zebra=LIGHT, first_col_bold=False, align_cols=None, bold_rows=(),
          cell_fills=None):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gt.table
    # remove estilo padrão (bandas azuis)
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r in range(nrows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_header = header and r == 0
            fill_color = header_fill if is_header else (zebra if (r % 2 == 0 and zebra is not None) else WHITE)
            if cell_fills and (r, c) in cell_fills:
                fill_color = cell_fills[(r, c)]
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill_color
            val = rows[r][c]
            runs = val if isinstance(val, list) else [(str(val), {})]
            al = PP_ALIGN.LEFT
            if align_cols and c in align_cols:
                al = align_cols[c]
            fill_text(cell.text_frame, [runs], size=size, color=header_color if is_header else INK,
                      bold=is_header or (first_col_bold and c == 0) or (r in bold_rows), align=al, space_after=0)
    return gt


# ---------- gráfico nativo ----------
def bar_chart(slide, x, y, w, h, categories, values, color=NAVY, number_format='#,##0', font_size=14,
              title=None, labels=None):
    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("Registros", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = title
        tp = ch.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size = Pt(font_size + 2)
        tp.runs[0].font.bold = True
        tp.runs[0].font.name = FONT
        tp.runs[0].font.color.rgb = NAVY
    ch.font.name = FONT
    ch.font.size = Pt(font_size)
    ch.font.color.rgb = INK
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(font_size)
    dl.font.name = FONT
    dl.font.color.rgb = INK
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = color
    if labels:  # rótulos de texto fixos (independentes do idioma/locale do PowerPoint)
        for i, lab in enumerate(labels):
            pdl = ser.points[i].data_label
            pdl.position = XL_LABEL_POSITION.OUTSIDE_END
            tf = pdl.text_frame
            tf.text = lab
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)
                    r.font.name = FONT
                    r.font.color.rgb = INK
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = max(values) * 1.22
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = rgb("C9CFDA")
    ca.tick_labels.font.size = Pt(font_size)
    ca.tick_labels.font.name = FONT
    ca.reverse_order = True  # primeira categoria no topo
    return gf


# ---------- moldura do modelo ----------
class Deck:
    def __init__(self, template_path):
        self.prs = Presentation(str(template_path))
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "Blank")
        self.bg = ASSETS / "modelo" / "tpl_bg.png"
        self.logo = ASSETS / "modelo" / "tpl_semica_logo.png"
        self.slide_count = 0

    def remove_slides(self, indexes):
        sldIdLst = self.prs.slides._sldIdLst
        ids = list(sldIdLst)
        for i in sorted(indexes, reverse=True):
            el = ids[i]
            self.prs.part.drop_rel(el.rId)
            sldIdLst.remove(el)
        # renumera os slides restantes (evita colisão de partname ao adicionar novos slides)
        from pptx.opc.packuri import PackURI
        for n, slide in enumerate(self.prs.slides, start=1):
            slide.part.partname = PackURI(f"/ppt/slides/slide{n}.xml")

    def move_slide_to_end(self, index):
        sldIdLst = self.prs.slides._sldIdLst
        el = list(sldIdLst)[index]
        sldIdLst.remove(el)
        sldIdLst.append(el)

    def content_slide(self, section, subtitle=None, track=None):
        s = self.prs.slides.add_slide(self.blank)
        # moldura idêntica ao modelo: fundo com faixa azul, retângulo branco e logo SEMICA
        bgp = s.shapes.add_picture(str(self.bg), Emu(20), Emu(1923), Emu(18287980), Emu(10285077))
        bgp.name = "Plano de fundo do modelo"
        wr = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(15888345), Emu(279974), Emu(2362200), Emu(863025))
        _strip_style(wr)
        wr.fill.solid(); wr.fill.fore_color.rgb = WHITE; wr.line.fill.background()
        wr.name = "Retângulo branco (modelo)"
        lg = s.shapes.add_picture(str(self.logo), Emu(15544800), Emu(279974), Emu(2512314), Emu(1234862))
        lg.name = "Logo SEMICA"
        t = s.shapes.add_textbox(Emu(1143000), Emu(558224), Inches(11.0), Emu(584775))
        t.name = "Título da seção"
        tf = t.text_frame
        tf.word_wrap = True
        fill_text(tf, [section], size=32, color=RGBColor(0, 0, 0), space_after=0)
        if subtitle:
            tb(s, 1.25, 1.2, 15.5, 0.55, [subtitle], size=24, color=NAVY, bold=True, space_after=0,
               name="Subtítulo")
        if track is not None:
            step_track(s, track)
        self.slide_count += 1
        return s

    @staticmethod
    def notes(slide, text):
        slide.notes_slide.notes_text_frame.text = text

    def save(self, path):
        self.prs.save(str(path))
