"""Gera a VERSÃO ENXUTA da apresentação do VIII SEMICA (mesmo modelo oficial e mesma identidade visual).

Uso:  python3 build_deck_enxuto.py  [saida.pptx]

Pedido da orientadora: slides mais limpos, tempo cronometrado e foco em como o trabalho foi feito e na resposta ao
objetivo. Cada slide tem uma mensagem (o subtítulo é a frase-chave) e um destaque visual; os detalhes foram para as notas
e para os slides de apoio, depois do encerramento, mostrados só se perguntarem. Capa, referências, agradecimentos,
encerramento e slides de apoio reaproveitam as funções de build_deck.py.
Também grava ROTEIRO_VERSAO_ENXUTA.md: tempo por slide, relógio acumulado e a fala de cada slide.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import build_deck as completo
from build_deck import FIG, HL, IMG, TEMPLATE, Y0, YB, YBT
from deck_lib import (GAP, GOLD, GOLD_TINT, GRAY, GREEN, GREEN_TINT, INK, L, LIGHT, NAVY, NAVY_TINT, PAD, R, RED,
                      RED_TINT, TEAL, TEAL_TINT, WHITE, Build, Deck, add_transition, circle, cols, icon, icon_badge,
                      line, pic, rect, rgb, table, tb)
from notas_fala_enxuta import APOIO, NOTAS

HERE = Path(__file__).resolve().parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "SEMICA_Fumaca_Cirurgica_Baccaro_ENXUTA.pptx"
ROTEIRO_MD = HERE / "ROTEIRO_VERSAO_ENXUTA.md"
DARK_GOLD = completo.DARK_GOLD
EMU = 914400


def bloco(s, x, y, w, h, title, body, icon_name, badge_fill=NAVY, fill=LIGHT, title_size=21, body_size=20,
          bullet=None, space_after=8, badge_d=0.8):
    """Card com ícone e título no topo e o texto centralizado na altura restante (sem vazio embaixo)."""
    rect(s, x, y, w, h, fill=fill)
    icon_badge(s, icon_name, x + PAD + badge_d / 2, y + PAD + badge_d / 2, d=badge_d, fill=badge_fill)
    tb(s, x + PAD + badge_d + 0.25, y + PAD, w - 2 * PAD - badge_d - 0.25, badge_d, [title], size=title_size,
       color=NAVY, bold=True, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    top = y + PAD + badge_d + 0.15
    tb(s, x + PAD, top, w - 2 * PAD, y + h - PAD - top, body, size=body_size, color=INK, bullet=bullet,
       space_after=space_after, anchor=MSO_ANCHOR.MIDDLE)


# =====================================================================
# 2 — INTRODUÇÃO
# =====================================================================
def slide_introducao(deck):
    s = deck.content_slide("INTRODUÇÃO", "Fumaça cirúrgica: exposição diária da equipe, efeitos ainda incertos")
    LW = 8.0
    foto = pic(s, IMG / "centro_cirurgico_fumaca.png", L, Y0, w=LW, name="Foto centro cirúrgico")
    fy = Y0 + foto.height / EMU
    tb(s, L, fy + 0.08, LW, 0.35, ["Imagem ilustrativa (gerada por IA): fumaça do eletrocautério no campo cirúrgico."],
       size=14, color=GRAY, space_after=0)
    BAND_H = 1.45
    BAND_Y = YB - BAND_H
    SY = fy + 0.6
    SH = BAND_Y - GAP - SY
    rect(s, L, SY, LW, SH, fill=NAVY_TINT)
    tb(s, L + PAD, SY + 0.18, LW - 2 * PAD, 0.4, ["Sintomas relatados pela equipe"], size=17, color=NAVY, bold=True,
       space_after=0)
    for (x, w), (ic, t) in zip(cols(4, L + PAD, L + LW - PAD, gap=0.1),
                               [("eye", "Olhos"), ("lungs", "Vias aéreas"), ("nose", "Mucosas"), ("exclaim", "Cefaleia")]):
        icon(s, ic, x + (w - 0.5) / 2, SY + 0.72, 0.5, color="navy")
        tb(s, x, SY + 1.27, w, 0.4, [t], size=17, color=INK, align=PP_ALIGN.CENTER, space_after=0)
    RX = L + LW + 0.6
    RW = R - RX
    rows = [
        ("smoke", "O que é", [[("Subproduto do ", {}), ("eletrocautério, do laser e do bisturi ultrassônico", {"bold": True}),
                              (" ao cortar e coagular tecido.", {})]]),
        ("flask", "O que contém", [[("Benzeno", {"bold": True}), (" e outros compostos voláteis, ", {}),
                                   ("partículas ultrafinas", {"bold": True}), (" e DNA viral, como o do HPV.", {})]]),
        ("warning", "O que se sabe", [[("Sintomas agudos frequentes; efeitos crônicos plausíveis, mas ", {}),
                                       ("sem relação causal estabelecida", {"bold": True, **HL}), (".", {})]]),
    ]
    row_h = (BAND_Y - GAP - Y0) / len(rows)
    for i, (ic, t, body) in enumerate(rows):
        yy = Y0 + i * row_h
        icon_badge(s, ic, RX + 0.5, yy + 0.5, d=1.0, fill=NAVY)
        tb(s, RX + 1.3, yy + 0.02, RW - 1.3, 0.5, [t], size=24, color=NAVY, bold=True, space_after=0)
        tb(s, RX + 1.3, yy + 0.6, RW - 1.3, row_h - 0.7, body, size=21, color=INK, space_after=0, line_spacing=1.02)
    rect(s, L, BAND_Y, R - L, BAND_H, fill=NAVY)
    icon(s, "question", L + PAD, BAND_Y + (BAND_H - 0.7) / 2, 0.7, color="gold")
    tb(s, L + PAD + 1.0, BAND_Y + 0.15, R - L - 2 * PAD - 1.0, BAND_H - 0.3,
       [[("Lacuna: ", {"bold": True, "color": GOLD}), ("as revisões anteriores eram descritivas; ", {}),
         ("faltava combinar os estudos numa estimativa e graduar a certeza da evidência.", {"bold": True})]],
       size=21, color=WHITE, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    return s


# =====================================================================
# 3 — OBJETIVO (os itens 1, 2 e 3 voltam na conclusão)
# =====================================================================
OBJETIVOS = [
    ("Prevalência de sintomas", "Metanálise", "ocular (principal) e respiratório (exploratório)"),
    ("Desfechos objetivos e biomarcadores", "Síntese narrativa", "função pulmonar, urina, mucosa nasal e HPV"),
    ("Certeza da evidência", "GRADE", "por desfecho, em cinco domínios"),
]


def slide_objetivos(deck):
    s = deck.content_slide("OBJETIVOS", "Pergunta: o que a fumaça cirúrgica faz à saúde da equipe?")
    OB_H = 1.55
    rect(s, L, Y0, R - L, OB_H, fill=NAVY)
    tb(s, L + 0.45, Y0 + 0.15, R - L - 0.9, OB_H - 0.3,
       [[("Objetivo geral: ", {"bold": True, "color": GOLD}),
         ("avaliar os desfechos de saúde associados à exposição ocupacional à fumaça cirúrgica na equipe de centro "
          "cirúrgico.", {})]], size=24, color=WHITE, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    CY, CH = Y0 + OB_H + GAP + 0.1, 3.45
    for i, ((x, w), (t, como, oque)) in enumerate(zip(cols(3), OBJETIVOS), start=1):
        rect(s, x, CY, w, CH, fill=LIGHT)
        circle(s, x + PAD + 0.45, CY + PAD + 0.45, 0.9, fill=NAVY, text=str(i), size=30)
        tb(s, x + PAD + 1.15, CY + PAD, w - 2 * PAD - 1.15, 0.9, [t], size=22, color=NAVY, bold=True, space_after=0,
           anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + PAD, CY + PAD + 1.05, w - 2 * PAD, CH - 2 * PAD - 1.05,
           [{"runs": [(como, {"bold": True, "color": RED, "size": 26})], "space_after": 6}, oque],
           size=21, color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    PL_Y = CY + CH + 0.35
    tb(s, L, PL_Y, R - L, 0.4, [[("Pergunta estruturada (PECO)", {"bold": True, "color": NAVY})]], size=18, space_after=0)
    peco = [("P", "População", "Equipe de centro cirúrgico"), ("E", "Exposição", "Fumaça de dispositivos de energia"),
            ("C", "Comparador", "Não expostos ou protegidos"), ("O", "Desfecho", "Efeitos à saúde")]
    PY = PL_Y + 0.5
    for (x, w), (letter, name, val) in zip(cols(4), peco):
        rect(s, x, PY, w, YB - PY, fill=TEAL_TINT)
        circle(s, x + PAD + 0.45, PY + (YB - PY) / 2, 0.9, fill=TEAL, text=letter, size=28)
        tb(s, x + PAD + 1.15, PY + 0.15, w - 2 * PAD - 1.15, YB - PY - 0.3,
           [{"runs": [(name, {"bold": True, "color": TEAL, "size": 20})], "space_after": 2}, val], size=19, color=INK,
           space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    return s


# =====================================================================
# 4 — MÉTODO EM 7 PASSOS
# =====================================================================
def slide_metodo(deck):
    s = deck.content_slide("METODOLOGIA", "Revisão sistemática em 7 passos, conforme o PRISMA 2020")
    b = Build(s)
    steps = [
        ("Protocolo", "PROSPERO", "registrado antes da extração"),
        ("Busca", "7 bases", "+ literatura cinzenta e citação, sem filtro de idioma"),
        ("Seleção", "10.910", "títulos e resumos lidos, um a um"),
        ("Extração", "Planilha", "padronizada: eventos e denominador"),
        ("Risco de viés", "JBI", "ferramenta para cada desenho de estudo"),
        ("Síntese", "Metanálise", "de prevalência, com efeitos aleatórios"),
        ("Certeza", "GRADE", "por desfecho, em cinco domínios"),
    ]
    sc = cols(len(steps), L, R, gap=0.2)
    cy = 2.6
    line(s, sc[0][0] + sc[0][1] / 2, cy, sc[-1][0] + sc[-1][1] / 2, cy, color=rgb("C9CFDA"), width=4)
    BX_Y, BX_H = 3.95, 2.5
    for i, ((x, w), (t, big, small)) in enumerate(zip(sc, steps), start=1):
        destaque = i == 6
        with b.step():
            circle(s, x + w / 2, cy, 1.0, fill=RED if destaque else NAVY, text=str(i), size=28)
            tb(s, x, 3.3, w, 0.5, [t], size=20, color=NAVY, bold=True, align=PP_ALIGN.CENTER, space_after=0)
            rect(s, x, BX_Y, w, BX_H, fill=RED_TINT if destaque else LIGHT)
            tb(s, x + 0.12, BX_Y + 0.22, w - 0.24, 0.6, [big], size=28 if i == 3 else 24, color=RED if destaque else NAVY,
               bold=True, align=PP_ALIGN.CENTER, space_after=0)
            tb(s, x + 0.15, BX_Y + 0.9, w - 0.3, BX_H - 1.05, [small], size=18, color=INK, align=PP_ALIGN.CENTER,
               space_after=0)
    SL_Y = BX_Y + BX_H + GAP + 0.05
    SL_H = 1.95
    with b.step():
        rect(s, L, SL_Y, R - L, SL_H, fill=NAVY)
        selos = [("clipboard", "Protocolo registrado", "PROSPERO CRD420251238727"),
                 ("check", "Relato completo", "checklist PRISMA 2020, 27 itens"),
                 ("table", "Ciência aberta", "código, planilhas e figuras no OSF")]
        for (x, w), (ic, t, d) in zip(cols(3, L + PAD, R - PAD, gap=0.4), selos):
            icon_badge(s, ic, x + 0.45, SL_Y + SL_H / 2, d=0.9, fill=GOLD)
            tb(s, x + 1.05, SL_Y + 0.3, w - 1.05, SL_H - 0.6,
               [{"runs": [(t, {"bold": True, "color": GOLD, "size": 22})], "space_after": 3}, d], size=19,
               color=WHITE, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    TL_Y = SL_Y + SL_H + GAP
    with b.step():
        for (x, w), body in zip(cols(2), [
            [[("Ferramentas: ", {"bold": True, "color": NAVY}), ("Zotero · Rayyan · OpenAlex · Python", {})]],
            [[("IA: ", {"bold": True, "color": NAVY}), ("apoio operacional, declarado; ", {}),
              ("as decisões foram humanas.", {"bold": True})]],
        ]):
            rect(s, x, TL_Y, w, YB - TL_Y, fill=NAVY_TINT)
            tb(s, x + PAD, TL_Y + 0.1, w - 2 * PAD, YB - TL_Y - 0.2, body, size=19, color=INK, space_after=0,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    b.apply()
    return s


# =====================================================================
# 5 — BUSCA E SELEÇÃO (passos 2 e 3)
# =====================================================================
def slide_selecao(deck):
    s = deck.content_slide("METODOLOGIA", "Passos 2 e 3 · Da busca ampla aos 22 relatórios incluídos", track={2, 3})
    b = Build(s)
    flow = [("16.757", "registros identificados", NAVY), ("10.910", "títulos e resumos lidos", NAVY),
            ("32", "buscados na íntegra", NAVY), ("26", "avaliados na íntegra", NAVY), ("22", "relatórios incluídos", RED)]
    between = [["− 5.839", "duplicatas", "− 8 não", "exportados"], ["− 10.882", "excluídos", "+ 4 por", "citação"],
               ["− 6 não", "recuperados"], ["− 4", "excluídos"]]
    FB_H = 2.9
    bw, bg = 2.6, (R - L - 5 * 2.6) / 4
    for i, (num, lab, col) in enumerate(flow):
        x = L + i * (bw + bg)
        with b.step():
            if i > 0:
                px = x - bg
                line(s, px + 0.08, Y0 + 0.95, x - 0.08, Y0 + 0.95, color=NAVY, width=2.5, arrow=True)
                tb(s, px + 0.02, Y0 + 1.15, bg - 0.04, 1.4, between[i - 1], size=16, color=GRAY, align=PP_ALIGN.CENTER,
                   space_after=0, line_spacing=0.95)
            rect(s, x, Y0, bw, FB_H, fill=RED_TINT if col == RED else LIGHT)
            tb(s, x + 0.12, Y0 + 0.28, bw - 0.24, 0.9, [num], size=50, color=col, bold=True, align=PP_ALIGN.CENTER,
               space_after=0)
            tb(s, x + 0.2, Y0 + 1.35, bw - 0.4, FB_H - 1.55, [lab], size=20, color=INK, align=PP_ALIGN.CENTER,
               space_after=0)
    CA_Y = Y0 + FB_H + 0.45
    CA_H = YBT - CA_Y
    (c1x, c1w), (c2x, c2w), (c3x, c3w) = cols(3)
    with b.step():
        bloco(s, c1x, CA_Y, c1w, CA_H, "Busca de alta sensibilidade", [
            [("7 bases", {"bold": True}), (", literatura cinzenta e busca por citação", {})],
            [("Sem filtro de idioma ou data", {"bold": True})],
            "Última busca: 09/07/2026",
        ], "search", bullet=True, body_size=21, space_after=10)
    with b.step():
        bloco(s, c2x, CA_Y, c2w, CA_H, "Ninguém contado duas vezes", [
            "7 publicações do mesmo grupo, com amostras sobrepostas",
            [("1 representante por análise", {"bold": True}), (", definido antes da síntese", {})],
        ], "people", bullet=True, body_size=21, space_after=10)
    with b.step():
        rect(s, c3x, CA_Y, c3w, CA_H, fill=RED_TINT)
        tb(s, c3x + PAD, CA_Y + PAD, c3w - 2 * PAD, CA_H - 2 * PAD, [
            {"runs": [("6 | 7", {"bold": True, "color": RED, "size": 54})], "space_after": 4},
            {"runs": [("amostras independentes nas metanálises", {"bold": True})], "space_after": 4},
            "ocular: n = 756 · respiratório: n = 1.428",
            {"runs": [("9 países · 19 transversais + 3 coortes", {"color": GRAY})], "space_before": 10},
        ], size=20, color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    b.apply()
    return s


# =====================================================================
# 6 — COMO OS ESTUDOS FORAM COMBINADOS E AVALIADOS (passos 5 a 7)
# =====================================================================
def slide_metanalise(deck):
    s = deck.content_slide("METODOLOGIA", "Passos 5 a 7 · Como os estudos foram avaliados e combinados", track={5, 6, 7})
    b = Build(s)
    pipe = [
        ("Proporção de cada estudo", ["33 / 81 = 40,7%"], "eventos ÷ participantes (ex.: İlçe et al., 2017)"),
        ("Escala logito", ["ln [ p ÷ (1 − p) ]"], "mantém estimativa e IC entre 0 e 100%"),
        ("Efeitos aleatórios", ["Paule-Mandel", "+ Hartung-Knapp"],
         "a prevalência verdadeira varia entre hospitais; IC honesto com poucos estudos"),
        ("Heterogeneidade e predição", ["I²  ·  intervalo de predição"],
         "quanto os estudos diferem e o que esperar num novo hospital"),
    ]
    PC_H = 4.2
    for i, ((x, w), (t, formula, body)) in enumerate(zip(cols(len(pipe)), pipe), start=1):
        with b.step():
            if i > 1:
                line(s, x - GAP + 0.03, Y0 + 1.9, x - 0.03, Y0 + 1.9, color=NAVY, width=2.5, arrow=True)
            rect(s, x, Y0, w, PC_H, fill=LIGHT)
            circle(s, x + PAD + 0.375, Y0 + PAD + 0.375, 0.75, fill=RED, text=str(i), size=24)
            tb(s, x + PAD + 0.95, Y0 + PAD - 0.05, w - 2 * PAD - 0.95, 0.85, [t], size=21, color=NAVY, bold=True,
               space_after=0, anchor=MSO_ANCHOR.MIDDLE)
            rect(s, x + 0.25, Y0 + 1.4, w - 0.5, 1.0, fill=WHITE, line=rgb("C9CFDA"))
            tb(s, x + 0.3, Y0 + 1.4, w - 0.6, 1.0, formula, size=22 if len(formula) == 1 else 19, color=RED, bold=True,
               align=PP_ALIGN.CENTER, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
            tb(s, x + PAD, Y0 + 2.55, w - 2 * PAD, PC_H - 2.55 - PAD, [body], size=20, color=INK, space_after=0,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    PN_Y = Y0 + PC_H + GAP + 0.05
    paineis = [
        ("balance", "Passo 5 · Risco de viés", NAVY, NAVY_TINT,
         [[("Ferramentas JBI", {"bold": True}), (" por desenho: prevalência (9 itens), transversal (8) e coorte (11)", {})]]),
        ("check", "Robustez da metanálise", TEAL, TEAL_TINT,
         [[("3 modelos", {"bold": True}), (", leave-one-out, Baujat e subgrupo por país", {})]]),
        ("target", "Passo 7 · Certeza (GRADE)", DARK_GOLD, GOLD_TINT,
         [[("5 domínios por desfecho", {"bold": True}), ("; parte de certeza baixa (estudos observacionais)", {})]]),
    ]
    for (x, w), (ic, t, badge, tint, body) in zip(cols(3), paineis):
        with b.step():
            bloco(s, x, PN_Y, w, YBT - PN_Y, t, body, ic, badge_fill=badge, fill=tint, body_size=21)
    b.apply()
    return s


# =====================================================================
# 7 — RESULTADO PRINCIPAL: SINTOMA OCULAR
# =====================================================================
def slide_ocular(deck):
    s = deck.content_slide("RESULTADOS", "Desfecho principal · Cerca de 4 em cada 10 profissionais relataram sintoma ocular")
    b = Build(s)
    FW = 12.3
    BAND_Y = 8.95
    with b.step():
        fp = pic(s, FIG / "fig2_forest_ocular_slide.png", L, Y0, w=FW, name="Figura 2 forest plot ocular")
        fb = Y0 + fp.height / EMU
        tb(s, L, fb + 0.08, FW, 0.6, [[("Figura 2. ", {"bold": True}),
                                       ("Quadrado = estudo (tamanho ∝ peso) · losango = estimativa combinada · "
                                        "linha vermelha = intervalo de predição.", {})]], size=14, color=GRAY, space_after=0)
    SRX = L + FW + GAP
    SW = R - SRX
    P1_H = 3.75
    with b.step():
        rect(s, SRX, Y0, SW, P1_H, fill=RED_TINT)
        tb(s, SRX + PAD, Y0 + 0.2, SW - 2 * PAD, 1.35, ["37,6%"], size=80, color=RED, bold=True, space_after=0)
        tb(s, SRX + PAD, Y0 + 1.6, SW - 2 * PAD, 0.4, ["prevalência combinada"], size=19, color=GRAY, space_after=0)
        tb(s, SRX + PAD, Y0 + 2.15, SW - 2 * PAD, P1_H - 2.15 - PAD, [
            [("IC95% ", {"bold": True}), ("21,9–56,5", {})],
            [("6 estudos · 756 profissionais", {"bold": True})],
        ], size=20, color=INK, space_after=4)
    P2_Y = Y0 + P1_H + GAP
    with b.step():
        rect(s, SRX, P2_Y, SW, YB - P2_Y, fill=GOLD_TINT)
        tb(s, SRX + PAD, P2_Y + PAD, SW - 2 * PAD, YB - P2_Y - 2 * PAD, [
            {"runs": [("Intervalo de predição", {"bold": True, "color": NAVY, "size": 22})], "space_after": 2},
            {"runs": [("7–83%", {"bold": True, "color": RED, "size": 66})], "space_after": 6},
            [("o que esperar num novo hospital: ", {}), ("depende do contexto", {"bold": True, **HL})],
        ], size=20, color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    with b.step():
        rect(s, L, BAND_Y, FW, YB - BAND_Y, fill=NAVY_TINT)
        tb(s, L + PAD, BAND_Y + 0.15, FW - 2 * PAD, YB - BAND_Y - 0.3,
           [[("Como ler: ", {"bold": True, "color": NAVY}), ("os estudos vão de 19% (Malásia) a 56% (Turquia); ", {}),
             ("I² = 93%: a variação entre contextos é real.", {"bold": True})]],
           size=19, color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    b.apply()
    return s


# =====================================================================
# 8 — SINTOMA RESPIRATÓRIO E ROBUSTEZ
# =====================================================================
def slide_respiratorio(deck):
    s = deck.content_slide("RESULTADOS", "Desfecho exploratório · Sintoma respiratório em 41,7%, estável entre modelos")
    FW = 10.2
    fp = pic(s, FIG / "fig3_forest_resp_slide.png", L, Y0, w=FW, name="Figura 3 forest plot respiratório")
    fb = Y0 + fp.height / EMU
    tb(s, L, fb + 0.08, FW, 0.6, [[("Figura 3. ", {"bold": True}),
                                   ("Sete amostras independentes; definições variam (tosse, ardência de garganta, vias "
                                    "aéreas superiores).", {})]], size=14, color=GRAY, space_after=0)
    DF_Y = 8.55
    rect(s, L, DF_Y, FW, YB - DF_Y, fill=GOLD_TINT)
    tb(s, L + PAD, DF_Y + 0.15, FW - 2 * PAD, YB - DF_Y - 0.3, [
        [("A definição importa: ", {"bold": True, "color": NAVY}), ("só com definições específicas, o I² respiratório cai "
                                                                    "de ", {}), ("91% para 61%", {"bold": True, **HL})],
    ], size=20, color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    RX = L + FW + GAP
    RW = R - RX
    P1_H = 2.75
    rect(s, RX, Y0, RW, P1_H, fill=RED_TINT)
    tb(s, RX + PAD, Y0 + 0.15, RW - 2 * PAD, 1.15, ["41,7%"], size=66, color=RED, bold=True, space_after=0)
    tb(s, RX + PAD, Y0 + 1.35, RW - 2 * PAD, P1_H - 1.35 - PAD, [
        {"runs": [("prevalência combinada", {"color": GRAY})], "space_after": 4},
        [("IC95% ", {"bold": True}), ("30,4–53,9  ·  ", {}), ("IP ", {"bold": True}), ("15–74%", {})],
        [("7 estudos · 1.428 profissionais", {"bold": True})],
    ], size=18, color=INK, space_after=2)
    RB_Y = Y0 + P1_H + GAP
    rect(s, RX, RB_Y, RW, YB - RB_Y, fill=LIGHT)
    tb(s, RX + PAD, RB_Y + PAD - 0.05, RW - 2 * PAD, 0.45, ["Muda o modelo, não muda o resultado"], size=21, color=NAVY,
       bold=True, space_after=0)
    TW = RW - 2 * PAD
    table(s, RX + PAD, RB_Y + 0.95, TW, [
        ["Análise", "Ocular", "Respiratório"],
        ["Logito + Hartung-Knapp", "37,6%", "41,7%"],
        ["Freeman-Tukey", "38,1%", "41,7%"],
        ["GLMM logístico", "37,5%", "42,0%"],
        ["Leave-one-out (faixa)", "34,3–42,1%", "38,9–45,5%"],
        ["Sem o resumo de congresso", "42,1%", "42,4%"],
    ], col_widths=[TW - 3.6, 1.7, 1.9], size=17, row_h=0.56, align_cols={1: PP_ALIGN.CENTER, 2: PP_ALIGN.CENTER},
          bold_rows=(1,))
    tb(s, RX + PAD, YB - PAD - 0.45, RW - 2 * PAD, 0.45, ["Modelo principal em negrito."], size=14, color=GRAY,
       space_after=0)
    return s


# =====================================================================
# 9 — DESFECHOS OBJETIVOS E BIOMARCADORES
# =====================================================================
def slide_objetivos_bio(deck):
    s = deck.content_slide("RESULTADOS", "Medidas objetivas apontam na mesma direção: alterações além do incômodo agudo")
    b = Build(s)
    tb(s, L, Y0, R - L, 0.4, [[("Expostos vs não expostos", {"bold": True, "color": NAVY}),
                               ("  ·  um estudo por desfecho", {"color": GRAY})]], size=18, space_after=0)
    tiles = [
        ("lungs", "Função pulmonar", "2,78 vs 3,14 L", "VEF1 (mediana); menor com mais tempo de centro cirúrgico",
         "Ibrahim et al., 2025 · Sudão · 92 vs 92"),
        ("vial", "Benzeno na urina", "202,9 vs 46,1", "µg/L de ácido t,t-mucônico, metabólito do benzeno",
         "Canicoba, 2025 · Brasil · 52 vs 49"),
        ("biotech", "Mucosa nasal", "70% vs 5%", "metaplasia ou hiperplasia na biópsia · RR 13,9",
         "Navarro et al., 2016 · México · 23 vs 20"),
        ("virus", "HPV na mucosa nasal", "8,96% vs 1,73%", "DNA de HPV detectado · RR 5,2",
         "Hu et al., 2021 · China · 469 vs 231"),
    ]
    TY = Y0 + 0.55
    BAND_H = 2.0
    TH = YB - BAND_H - GAP - TY
    for (x, w), (ic, t, big, what, src) in zip(cols(4), tiles):
        with b.step():
            rect(s, x, TY, w, TH, fill=LIGHT)
            icon_badge(s, ic, x + w / 2, TY + PAD + 0.55, d=1.1, fill=NAVY)
            tb(s, x + PAD, TY + 1.65, w - 2 * PAD, 0.5, [t], size=22, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
               space_after=0)
            tb(s, x + 0.12, TY + 2.3, w - 0.24, 0.75, [big], size=36, color=RED, bold=True, align=PP_ALIGN.CENTER,
               space_after=0)
            tb(s, x + PAD, TY + 3.15, w - 2 * PAD, 1.05, [what], size=19, color=INK, align=PP_ALIGN.CENTER,
               space_after=0)
            tb(s, x + PAD, TY + TH - PAD - 0.6, w - 2 * PAD, 0.6, [src], size=15, color=GRAY, align=PP_ALIGN.CENTER,
               space_after=0, anchor=MSO_ANCHOR.BOTTOM)
    BY = YB - BAND_H
    with b.step():
        rect(s, L, BY, R - L, BAND_H, fill=NAVY)
        for (x, w), (num, lab) in zip(cols(3, L + PAD, R - PAD, gap=0.5), [
            ("1", "estudo por desfecho: sinais, não prova de doença"),
            (None, "certeza muito baixa em todos os desfechos (GRADE)"),
            ("0", "estudos elegíveis estimaram câncer ou DPOC"),
        ]):
            if num is None:  # escala do GRADE: 1 de 4 círculos preenchido = muito baixa
                nw = 2.0
                for j in range(4):
                    circle(s, x + 0.25 + j * 0.5, BY + BAND_H / 2, 0.38, fill=GOLD if j == 0 else NAVY, line=GOLD)
            else:
                nw = 0.8
                tb(s, x, BY + 0.2, nw, BAND_H - 0.4, [num], size=44, color=GOLD, bold=True, space_after=0,
                   anchor=MSO_ANCHOR.MIDDLE)
            tb(s, x + nw + 0.15, BY + 0.2, w - nw - 0.15, BAND_H - 0.4, [lab], size=20, color=WHITE, space_after=0,
               anchor=MSO_ANCHOR.MIDDLE)
    b.apply()
    return s


# =====================================================================
# 10 — DISCUSSÃO
# =====================================================================
def slide_discussao(deck):
    s = deck.content_slide("DISCUSSÃO", "O que o trabalho acrescenta, como ler os resultados e seus limites")
    BAND_H = 1.6
    CH = YB - BAND_H - GAP - Y0
    colunas = [
        ("thumbsup", "O que acrescenta", GREEN, GREEN_TINT, [
            "Estimativa combinada, com intervalo de predição",
            "Certeza graduada (GRADE) por desfecho",
            "Protocolo registrado e análise aberta no OSF"]),
        ("compass", "Como ler", NAVY, NAVY_TINT, [
            [("Prevalência em expostos: ", {"bold": True}), ("carga de sintomas, não efeito causal", {})],
            [("I² > 90% é esperado: ", {"bold": True}), ("ler pelo intervalo de predição", {})],
            "Marcadores objetivos não são doença crônica"]),
        ("warning", "Limitações (declaradas)", RED, RED_TINT, [
            "Estudos transversais e sintomas autorrelatados",
            "Sem grupo não exposto nas metanálises",
            "Revisor único; 6 relatórios não recuperados"]),
    ]
    for (x, w), (ic, t, badge, tint, body) in zip(cols(3), colunas):
        bloco(s, x, Y0, w, CH, t, body, ic, badge_fill=badge, fill=tint, title_size=23, body_size=22, bullet=True,
              space_after=26, badge_d=0.9)
    rect(s, L, YB - BAND_H, R - L, BAND_H, fill=NAVY)
    icon(s, "question", L + PAD, YB - BAND_H + (BAND_H - 0.7) / 2, 0.7, color="gold")
    tb(s, L + PAD + 1.0, YB - BAND_H + 0.15, R - L - 2 * PAD - 1.0, BAND_H - 0.3,
       [[("Vazio de evidência: ", {"bold": True, "color": GOLD}), ("nenhum estudo elegível estimou câncer ou DPOC. ", {}),
         ("Ausência de estudos não é ausência de risco.", {"bold": True})]],
       size=21, color=WHITE, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    return s


# =====================================================================
# 11 — CONCLUSÃO: RESPOSTA AO OBJETIVO
# =====================================================================
def slide_conclusao(deck):
    s = deck.content_slide("CONCLUSÃO", "Resposta ao objetivo")
    b = Build(s)
    LW = 11.0
    BAND_H = 1.75
    TOP_B = YB - BAND_H - GAP
    respostas = [
        [("Frequentemente relatados: ", {}), ("ocular 37,6% e respiratório 41,7%", {"bold": True}),
         (", com grande variação entre contextos.", {})],
        [("Estudos isolados sugerem ", {}), ("alterações além do incômodo agudo", {"bold": True}),
         (", sem demonstrar doença crônica atribuível.", {})],
        [("Muito baixa", {"bold": True, "color": RED}), (" em todos os desfechos; nenhum estudo elegível sobre câncer "
                                                          "ou DPOC.", {})],
    ]
    rh = (TOP_B - Y0 - 2 * GAP) / 3
    for i, ((t, _, _), body) in enumerate(zip(OBJETIVOS, respostas)):
        y = Y0 + i * (rh + GAP)
        with b.step():
            rect(s, L, y, LW, rh, fill=LIGHT)
            circle(s, L + PAD + 0.45, y + rh / 2, 0.9, fill=NAVY, text=str(i + 1), size=30)
            tb(s, L + PAD + 1.2, y + 0.12, LW - 2 * PAD - 1.2, rh - 0.24,
               [{"runs": [(t, {"bold": True, "color": NAVY, "size": 21})], "space_after": 4}, body], size=19,
               color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    PX = L + LW + GAP
    PW = R - PX
    with b.step():
        ph = pic(s, IMG / "exaustao_n95.png", PX, Y0, w=PW, name="Foto exaustão e N95")
        py_ = Y0 + ph.height / EMU
        tb(s, PX, py_ + 0.06, PW, 0.35, ["Imagem ilustrativa (gerada por IA)."], size=14, color=GRAY, space_after=0)
        BX_Y = Y0 + 2 * (rh + GAP)   # alinhada com a resposta 3, ao lado
        rect(s, PX, BX_Y, PW, TOP_B - BX_Y, fill=NAVY_TINT)
        BX_H = TOP_B - BX_Y
        tb(s, PX + PAD, BX_Y + 0.2, PW - 2 * PAD, 0.4, ["Proteção recomendada (NIOSH, 1996)"], size=17, color=NAVY,
           bold=True, space_after=0)
        IY = BX_Y + 0.65 + (BX_H - 0.65 - 1.08) / 2   # ícone + rótulo centralizados abaixo do título
        for (x, w), (ic, t) in zip(cols(3, PX + PAD, R - PAD, gap=0.15),
                                   [("fan", "Exaustão na fonte"), ("shield", "N95/PFF2"), ("graduation", "Educação")]):
            icon_badge(s, ic, x + w / 2, IY + 0.3, d=0.6, fill=NAVY)
            tb(s, x, IY + 0.68, w, 0.4, [t], size=17, color=INK, align=PP_ALIGN.CENTER, space_after=0)
    with b.step():
        rect(s, L, YB - BAND_H, R - L, BAND_H, fill=NAVY)
        tb(s, L + 0.45, YB - BAND_H + 0.15, R - L - 0.9, BAND_H - 0.3,
           [[("Implicação: ", {"bold": True, "color": GOLD}), ("pela carga de sintomas e pelo princípio da precaução, ", {}),
             ("proteger a equipe parece justificado", {"bold": True}),
             ("; o risco de longo prazo depende de estudos longitudinais.", {})]],
           size=21, color=WHITE, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    b.apply()
    return s


# =====================================================================
# 12 — PERSPECTIVAS FUTURAS
# =====================================================================
def slide_perspectivas(deck):
    s = deck.content_slide("PERSPECTIVAS FUTURAS", "O que falta para responder ao risco de longo prazo")
    agenda = [("hourglass", "Coortes", "com medida objetiva da exposição (PM2,5, voláteis)"),
              ("list", "Definições padronizadas", "sintoma, recordatório e instrumentos validados"),
              ("lungs", "Desfechos crônicos", "incidência de câncer e DPOC"),
              ("fan", "Efetividade da proteção", "exaustão, N95/PFF2 e educação, antes e depois")]
    AH = 4.6
    for (x, w), (ic, t, body) in zip(cols(4), agenda):
        rect(s, x, Y0, w, AH, fill=LIGHT)
        icon_badge(s, ic, x + w / 2, Y0 + PAD + 0.6, d=1.2, fill=NAVY)
        tb(s, x + PAD, Y0 + 1.75, w - 2 * PAD, 0.9, [t], size=22, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
           space_after=0, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + PAD, Y0 + 2.7, w - 2 * PAD, AH - 2.7 - PAD, [body], size=20, color=INK, align=PP_ALIGN.CENTER,
           space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    PY = Y0 + AH + GAP
    rect(s, L, PY, R - L, YB - PY, fill=GOLD_TINT)
    tb(s, L + PAD, PY + PAD - 0.05, R - L - 2 * PAD, 0.45, ["Para este trabalho"], size=22, color=NAVY, bold=True,
       space_after=0)
    IT_Y = PY + 0.95
    IT_H = YB - PAD - IT_Y
    for (x, w), (ic, body) in zip(cols(3, L + PAD, R - PAD, gap=0.5), [
        ("upload", [("Submissão ", {"bold": True}), ("do manuscrito a periódico indexado", {})]),
        ("people", [("Triagem e extração ", {"bold": True}), ("com segundo revisor independente", {})]),
        ("calendar", [("Atualização da busca ", {"bold": True}), ("e dos 6 relatórios não recuperados", {})]),
    ]):
        icon(s, ic, x, IT_Y + (IT_H - 0.65) / 2, 0.65, color="navy")
        tb(s, x + 0.9, IT_Y, w - 0.9, IT_H, [body], size=20, color=INK, space_after=0, anchor=MSO_ANCHOR.MIDDLE)
    return s


# =====================================================================
# ordem, tempos e montagem
# =====================================================================
ROTEIRO = [("capa", completo.slide_capa, 15), ("introducao", slide_introducao, 45), ("objetivos", slide_objetivos, 25),
           ("metodo", slide_metodo, 50), ("selecao", slide_selecao, 35), ("metanalise", slide_metanalise, 50),
           ("ocular", slide_ocular, 55), ("respiratorio", slide_respiratorio, 35),
           ("objetivos_bio", slide_objetivos_bio, 45), ("discussao", slide_discussao, 40),
           ("conclusao", slide_conclusao, 35), ("perspectivas", slide_perspectivas, 15),
           ("referencias", completo.slide_referencias, 5), ("agradecimentos", completo.slide_agradecimentos, 10),
           ("encerramento", completo.slide_encerramento, None)]
# slides de apoio (depois do encerramento), na ordem citada nas notas do encerramento
APOIO_ORDEM = ["trajetoria", "passos12", "selecao", "passos45", "passo6", "passo7", "respiratorio", "objetivos_bio",
               "limitacoes"]


def mmss(x):
    return f"{x // 60}:{x % 60:02d}"


def gravar_roteiro(prs, n_principais):
    linhas, falas, fim = [], [], 0
    for i, sl in enumerate(prs.slides, start=1):
        if i > n_principais:
            break
        nota = sl.notes_slide.notes_text_frame.text
        m = re.match(r"\[(\d+):(\d+)–(\d+):(\d+)\]\s*(.+)", nota)
        fala = re.search(r"FALA:\n(.*?)(?:\n\s*\n|$)", nota, re.S)
        if not m:
            linhas.append(f"| {i} | Encerramento (fica na tela durante as perguntas) | — | — | — |")
            continue
        a, b = int(m.group(1)) * 60 + int(m.group(2)), int(m.group(3)) * 60 + int(m.group(4))
        fim = b
        secao = re.sub(r"\s*\(~\d+ s[^)]*\).*$", "", m.group(5)).split("←")[0].strip()
        pal = len(fala.group(1).split()) if fala else 0
        linhas.append(f"| {i} | {secao.capitalize()} | {b - a} s | {pal} | {mmss(a)} → {mmss(b)} |")
        falas.append(f"### {i}. {secao.capitalize()}  ·  {mmss(a)} → {mmss(b)}\n\n{fala.group(1).strip() if fala else ''}\n")
    total_pal = sum(int(ln.split("|")[4]) for ln in linhas if ln.split("|")[4].strip().isdigit())
    md = [
        "# Roteiro cronometrado — versão enxuta",
        "",
        f"Tempo-alvo da fala: **{mmss(fim)}** (limite do evento: 10 min), {total_pal} palavras "
        f"(~{round(total_pal / (fim / 60))} palavras por minuto, ritmo calmo, com folga para apontar a figura e respirar).",
        "",
        "Ensaie com cronômetro: se em algum marco do relógio você estiver mais de 15 s atrasado, corte o detalhe do "
        "slide seguinte e siga a frase-chave (o subtítulo). Os slides de apoio, depois do encerramento, ficam para as "
        "perguntas; o aprendizado (slide de apoio 16) só se a banca der a oportunidade.",
        "",
        "| Slide | Seção | Tempo | Palavras | Relógio |",
        "|---|---|---|---|---|",
        *linhas,
        "",
        "## Fala, slide a slide",
        "",
        *falas,
    ]
    ROTEIRO_MD.write_text("\n".join(md), encoding="utf-8")
    return fim, total_pal


def main(out=OUT):
    completo.make_qr()
    deck = Deck(TEMPLATE)
    for key, fn, dur in ROTEIRO:
        deck.notes(fn(deck), NOTAS[key], dur=dur)
    n_principais = len(deck.prs.slides)
    deck.section_prefix = "APOIO · "
    for key in APOIO_ORDEM:
        deck.notes(getattr(completo, "slide_" + key)(deck), APOIO[key])
    for sl in deck.prs.slides:
        add_transition(sl)
    deck.save(out)
    fim, pal = gravar_roteiro(deck.prs, n_principais)
    print(f"OK: {out} ({n_principais} slides principais + {len(APOIO_ORDEM)} de apoio; roteiro de {mmss(fim)}, "
          f"{pal} palavras) e {ROTEIRO_MD.name}")


if __name__ == "__main__":
    main()
