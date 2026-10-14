"""Auditoria de layout: mede o respiro real do texto dentro dos cards.

Usa a geometria do PPTX (python-pptx) e a posição real de cada palavra no PDF
renderizado pelo LibreOffice (pdftotext -bbox-layout). Liberation Sans tem as
mesmas métricas do Arial, então a quebra de linha é a mesma do PowerPoint.

Uso:  python3 auditar_layout.py deck.pptx deck.pdf [--min-pad 0.25]
"""
from __future__ import annotations

import re
import subprocess
import sys
from html import unescape
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

EMU_IN = 914400.0
SLIDE_W_IN = 20.0


def words_by_page(pdf: Path):
    xml = subprocess.run(["pdftotext", "-bbox-layout", str(pdf), "-"], capture_output=True, text=True,
                         check=True).stdout
    pages = []
    for pm in re.finditer(r'<page width="([\d.]+)" height="([\d.]+)">(.*?)</page>', xml, re.S):
        pw = float(pm.group(1))
        k = SLIDE_W_IN / pw  # pt -> polegadas do slide
        ws = []
        for lm in re.finditer(r"<line[^>]*>(.*?)</line>", pm.group(3), re.S):
            for wm in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>',
                                  lm.group(1)):
                x0, y0, x1, y1 = (float(wm.group(i)) * k for i in range(1, 5))
                ws.append((x0, y0, x1, y1, unescape(wm.group(5))))
        pages.append(ws)
    return pages


def norm(t: str) -> str:
    return re.sub(r"[^\w%]+", "", t.lower())


def shape_box(sh):
    return (sh.left / EMU_IN, sh.top / EMU_IN, (sh.left + sh.width) / EMU_IN, (sh.top + sh.height) / EMU_IN)


def is_container(sh):
    if sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
        return False
    try:
        st = sh.auto_shape_type
    except Exception:
        return False
    if "RECTANGLE" not in str(st):
        return False
    if sh.name.startswith(("Retângulo branco", "Faixa zebrada")):
        return False
    try:
        if sh.fill.type is None:
            return False
    except Exception:
        return False
    x0, y0, x1, y1 = shape_box(sh)
    return (x1 - x0) * (y1 - y0) > 0.9 and (x1 - x0) < 19.5


def match_text(tokens, words, box):
    """Localiza no PDF as palavras de uma caixa de texto (casamento sequencial)."""
    bx0, by0, bx1, by1 = box
    cand = [w for w in words if w[0] >= bx0 - 0.15 and w[2] <= bx1 + 0.4 and w[1] >= by0 - 0.15 and w[1] <= by1 + 4.0]
    cand.sort(key=lambda w: (round(w[1], 1), w[0]))
    toks = [norm(t) for t in tokens if norm(t)]
    if not toks:
        return None
    best = None
    for start, w in enumerate(cand):
        if norm(w[4]) != toks[0] or w[1] > by1:
            continue
        got = [w]
        ti = 1
        for w2 in cand[start + 1:]:
            if ti >= len(toks):
                break
            if norm(w2[4]) == toks[ti]:
                got.append(w2)
                ti += 1
        if best is None or len(got) > len(best):
            best = got
        if ti >= len(toks):
            break
    if not best or len(best) < 0.85 * len(toks):
        return None
    return (min(w[0] for w in best), min(w[1] for w in best), max(w[2] for w in best), max(w[3] for w in best))


def main():
    pptx, pdf = Path(sys.argv[1]), Path(sys.argv[2])
    min_pad = 0.25
    if "--min-pad" in sys.argv:
        min_pad = float(sys.argv[sys.argv.index("--min-pad") + 1])
    prs = Presentation(str(pptx))
    pages = words_by_page(pdf)
    total = 0
    for si, (slide, words) in enumerate(zip(prs.slides, pages), start=1):
        issues = []
        conts = [(sh, shape_box(sh)) for sh in slide.shapes if is_container(sh)]
        texts = []
        for sh in slide.shapes:
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and "OVAL" in str(getattr(sh, "auto_shape_type", "")):
                continue
            if is_container(sh):
                continue
            box = shape_box(sh)
            rb = match_text(sh.text_frame.text.split(), words, box)
            if rb is None:
                continue
            texts.append((sh, box, rb))
        # respiro dentro dos containers
        cont_content = {id(c): [] for c, _ in conts}
        for sh, box, rb in texts:
            inside = [(c, cb) for c, cb in conts if cb[0] - 0.01 <= box[0] + 0.02 <= cb[2] and cb[1] - 0.01 <= box[1] + 0.02 <= cb[3]]
            if not inside:
                if rb[2] > 19.55 or rb[3] > 10.95:
                    issues.append(f"FORA DA ÁREA  '{sh.text_frame.text[:40]}'  fim x={rb[2]:.2f} y={rb[3]:.2f}")
                continue
            c, cb = min(inside, key=lambda t: (t[1][2] - t[1][0]) * (t[1][3] - t[1][1]))
            cont_content[id(c)].append(rb)
            pads = {"esq": rb[0] - cb[0], "dir": cb[2] - rb[2], "topo": rb[1] - cb[1], "base": cb[3] - rb[3]}
            lab = sh.text_frame.text.replace("\n", " ")[:45]
            if pads["base"] < 0.02:
                issues.append(f"TRANSBORDA    '{lab}'  passa {abs(pads['base']):.2f} in da base do card")
            for k, v in pads.items():
                lim = min_pad if k != "topo" else min_pad - 0.08
                if v < lim and not (k == "base" and pads["base"] < 0.02):
                    issues.append(f"APERTADO {k:<4} '{lab}'  respiro {v:.2f} in")
        # espaço vazio nos containers
        for c, cb in conts:
            items = cont_content[id(c)]
            if not items:
                continue
            h = cb[3] - cb[1]
            empty = cb[3] - max(r[3] for r in items)
            if empty > max(1.0, 0.33 * h):
                issues.append(f"VAZIO         card em y={cb[1]:.2f}–{cb[3]:.2f} x={cb[0]:.2f}: {empty:.2f} in livres na base")
        # sobreposição de textos renderizados
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i][2], texts[j][2]
                ix = min(a[2], b[2]) - max(a[0], b[0])
                iy = min(a[3], b[3]) - max(a[1], b[1])
                if ix > 0.05 and iy > 0.04:
                    issues.append(f"SOBREPOSIÇÃO  '{texts[i][0].text_frame.text[:25]}' × '{texts[j][0].text_frame.text[:25]}'")
        # alinhamento: containers na mesma faixa com bordas quase iguais
        for i in range(len(conts)):
            for j in range(i + 1, len(conts)):
                a, b = conts[i][1], conts[j][1]
                vo = min(a[3], b[3]) - max(a[1], b[1])
                if vo > 0.5 * min(a[3] - a[1], b[3] - b[1]) and (b[0] >= a[2] - 0.01 or a[0] >= b[2] - 0.01):
                    for k, nm in ((1, "topo"), (3, "base")):
                        d = abs(a[k] - b[k])
                        if 0.015 < d < 0.4:
                            issues.append(f"DESALINHADO   {nm} de cards lado a lado difere {d:.2f} in (x={a[0]:.2f} e x={b[0]:.2f})")
        if issues:
            print(f"\n=== Slide {si} ({len(issues)})")
            for it in issues:
                print("  ", it)
            total += len(issues)
    print(f"\nTOTAL de achados: {total}")


if __name__ == "__main__":
    main()
